from pathlib import Path
import json
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

root = Path(__file__).resolve().parents[2]
candidates = []
for path in root.rglob('*.pptx'):
    if path.name.startswith('~'):
        continue
    try:
        deck = Presentation(str(path))
        if len(deck.slides) == 27:
            candidates.append((len(path.parts), path, deck))
    except Exception:
        pass
_, source, prs = max(candidates, key=lambda item: item[0])
sw, sh = prs.slide_width, prs.slide_height
slides = []
for index, slide in enumerate(prs.slides, 1):
    shapes = []
    for shape_index, shape in enumerate(slide.shapes):
        if hasattr(shape, 'text') and shape.text.strip():
            continue
        kind = str(shape.shape_type)
        if shape.shape_type == MSO_SHAPE_TYPE.LINE:
            kind = 'line'
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            kind = 'image'
        elif shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            kind = 'table'
        elif shape.shape_type == MSO_SHAPE_TYPE.CHART:
            kind = 'chart'
        fill = None
        line = None
        try:
            if shape.fill.type and shape.fill.fore_color.rgb:
                fill = '#' + str(shape.fill.fore_color.rgb)
        except Exception:
            pass
        try:
            if shape.line.color.type and shape.line.color.rgb:
                line = '#' + str(shape.line.color.rgb)
        except Exception:
            pass
        shapes.append({'id': f'{index}-{shape_index}', 'type': kind, 'x': round(shape.left/sw*1280,1), 'y': round(shape.top/sh*720,1), 'w': round(shape.width/sw*1280,1), 'h': round(shape.height/sh*720,1), 'fill': fill, 'line': line})
    slides.append({'slide': index, 'shapes': shapes})
(root/'guanyu-html-builder/demo/ppt-shapes.json').write_text(json.dumps({'source':str(source),'slides':slides},ensure_ascii=False,indent=2),encoding='utf-8')
print('extracted shape data for',len(slides),'slides')
