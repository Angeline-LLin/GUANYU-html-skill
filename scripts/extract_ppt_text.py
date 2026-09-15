from pathlib import Path
import json
from pptx import Presentation

root = Path(__file__).resolve().parents[2]
candidates = []
for path in root.rglob('*.pptx'):
    if path.name.startswith('~'):
        continue
    try:
        deck = Presentation(str(path))
        candidates.append((len(deck.slides), path, deck))
    except Exception:
        continue
slide_count, source, prs = max(candidates, key=lambda item: (item[0], len(item[1].parts)))
slides = []
for index, slide in enumerate(prs.slides, 1):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text.strip():
            value = ' '.join(line.strip() for line in shape.text.splitlines() if line.strip())
            if value and value not in texts:
                texts.append(value)
    slides.append({'slide': index, 'texts': texts})
output = root / 'guanyu-html-builder' / 'demo' / 'ppt-content.json'
output.write_text(json.dumps({'source': str(source), 'slides': slides}, ensure_ascii=False, indent=2), encoding='utf-8')
print(output)
